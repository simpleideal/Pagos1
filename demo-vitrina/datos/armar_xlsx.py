#!/usr/bin/env python3
"""Arma la plantilla .xlsx para importar en Google Sheets."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
CSV = ROOT / "plantilla.csv"
XLSX = ROOT / "plantilla-google-sheets.xlsx"

COLS = [
    ("producto", "Nombre que ve el cliente. No lo dejes vacío."),
    ("precio_normal", "Precio de vitrina, solo números, sin $."),
    ("precio_oferta", "Precio de hoy. Si va vacío, se usa el normal."),
    ("hoy", "si o no. Solo si sale en la vitrina."),
    ("activo", "si o no. no = ni se publica."),
    ("nota", "Una línea. Puede ir vacía."),
    ("foto", "URL https://… o vacío. En Drive: cualquiera con el enlace."),
    ("unidades", "Número. 0 = Agotado. Vacío = no se muestra el cupo."),
    ("categoria", "Opcional. Agrupa: Horno, Para llevar…"),
]

ANCHO = {
    "producto": 28,
    "precio_normal": 16,
    "precio_oferta": 16,
    "hoy": 10,
    "activo": 10,
    "nota": 42,
    "foto": 28,
    "unidades": 12,
    "categoria": 16,
}

COMENTARIO = (
    "No cambies estos nombres.\n"
    "Esta pestaña tiene que llamarse ofertas."
)


def leer_filas():
    texto = CSV.read_text(encoding="utf-8")
    lineas = [ln for ln in texto.splitlines() if ln.strip()]
    encabezado = lineas[0].split(",")
    filas = []
    for ln in lineas[1:]:
        # CSV simple de la plantilla: las notas no llevan comas.
        filas.append(ln.split(","))
    return encabezado, filas


def estilo_cabecera(celda, relleno):
    celda.font = Font(bold=True, color="FFFFFFFF", name="Calibri", size=11)
    celda.fill = PatternFill("solid", fgColor=relleno)
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    celda.border = Border(
        left=Side(style="thin", color="FF1F3F28"),
        right=Side(style="thin", color="FF1F3F28"),
        top=Side(style="thin", color="FF1F3F28"),
        bottom=Side(style="thin", color="FF1F3F28"),
    )


def armar():
    encabezado, filas = leer_filas()
    assert encabezado == [c[0] for c in COLS], encabezado

    wb = Workbook()
    ofertas = wb.active
    ofertas.title = "ofertas"

    fill_base = "FF2F5D3A"
    fill_hoy = "FFC45C26"
    for i, (nombre, ayuda) in enumerate(COLS, start=1):
        celda = ofertas.cell(1, i, nombre)
        estilo_cabecera(celda, fill_hoy if nombre == "hoy" else fill_base)
        celda.comment = Comment(ayuda + "\n\n" + COMENTARIO, "Simple Ideal")
        ofertas.column_dimensions[get_column_letter(i)].width = ANCHO[nombre]

    for r, fila in enumerate(filas, start=2):
        for c, valor in enumerate(fila, start=1):
            nombre = encabezado[c - 1]
            celda = ofertas.cell(r, c)
            if nombre in ("precio_normal", "precio_oferta", "unidades") and valor != "":
                celda.value = int(valor)
                celda.number_format = "0"
            else:
                celda.value = valor
            if nombre == "hoy":
                celda.fill = PatternFill("solid", fgColor="FFF7E3C6")
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    ofertas.freeze_panes = "A2"
    ofertas.auto_filter.ref = f"A1:I{len(filas) + 1}"
    ofertas.row_dimensions[1].height = 28
    ofertas.sheet_properties.tabColor = "2F5D3A"

    si_no = DataValidation(
        type="list",
        formula1='"si,no"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Usa si o no",
        error="Escribe si o no, en minúscula.",
        promptTitle="Hoy / activo",
        prompt="si o no",
        showInputMessage=True,
    )
    ofertas.add_data_validation(si_no)
    si_no.add(f"D2:D200")
    si_no.add(f"E2:E200")

    uso = wb.create_sheet("como_usar")
    uso.sheet_properties.tabColor = "C45C26"
    uso.column_dimensions["A"].width = 88
    uso["A1"] = "Cómo usar esta hoja (el local no entra a GitHub)"
    uso["A1"].font = Font(bold=True, size=14, color="FF2F5D3A")
    lineas_uso = [
        "",
        "Esta pestaña no se publica. La vitrina lee solo la pestaña ofertas.",
        "",
        "Cada tarde:",
        "1. En lo que quieres liquidar: hoy = si. Completa precio_oferta y unidades.",
        "2. En lo demás: hoy = no.",
        "3. Si se acabó: unidades = 0 (sale Agotado).",
        "4. Si un producto no debe verse nunca: activo = no.",
        "",
        "No cambies la fila 1 ni el nombre de la pestaña ofertas.",
        "No pongas $ en los precios. Solo el número: 4500.",
        "foto puede ir vacía. Si pegas un enlace de Drive, que sea “cualquiera con el enlace”.",
        "",
        "Una vez (Simple Ideal): Compartir → cualquiera con el enlace puede ver.",
        "Copia el enlace de la hoja y pégalo en DATOS DE LA VITRINA → hoja.",
        "El llavero NFC y el QR siguen abriendo la misma URL.",
        "",
        "Para otro local: Archivo → Hacer una copia. Cada pastelería tiene su hoja.",
    ]
    for i, texto in enumerate(lineas_uso, start=2):
        uso.cell(i, 1, texto)
        uso.cell(i, 1).alignment = Alignment(wrap_text=True)
        uso.row_dimensions[i].height = 18

    wb.save(XLSX)
    print("escrito", XLSX)


if __name__ == "__main__":
    armar()
